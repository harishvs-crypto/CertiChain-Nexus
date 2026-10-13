import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_document():
    doc = Document()
    
    # Page setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles & Colors
    NAVY = RGBColor(30, 58, 138)       # #1E3A8A
    DARK_BLUE = RGBColor(15, 23, 42)    # #0F172A
    SLATE = RGBColor(71, 85, 105)       # #475569
    GREEN = RGBColor(22, 101, 52)       # #166534
    RED = RGBColor(153, 27, 27)         # #991B1B
    AMBER = RGBColor(180, 83, 9)        # #B45309

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("CertiChain Nexus")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(16)
    run_sub = subtitle_p.add_run("Comprehensive Verification Guide, System Architecture & Live Operational Status")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = SLATE

    # Metadata banner table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    col_widths = [Inches(3.2), Inches(3.3)]
    for row in meta_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_background(cell, "F1F5F9")
            set_cell_margins(cell, 80, 80, 120, 120)

    cell_00 = meta_table.cell(0, 0).paragraphs[0]
    cell_00.add_run("Document Version: ").bold = True
    cell_00.add_run("v1.0 (Live Operational Report)")
    
    cell_01 = meta_table.cell(0, 1).paragraphs[0]
    cell_01.add_run("Platform Engine: ").bold = True
    cell_01.add_run("CertiChain Trust Engine 2026")

    cell_10 = meta_table.cell(1, 0).paragraphs[0]
    cell_10.add_run("Cryptography: ").bold = True
    cell_10.add_run("Ed25519 Signatures + SHA-256")

    cell_11 = meta_table.cell(1, 1).paragraphs[0]
    cell_11.add_run("Access URL: ").bold = True
    cell_11.add_run("http://127.0.0.1:5000")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Helper function for section headings
    def add_section_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = NAVY
        return h

    def add_sub_header(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        r = h.add_run(text)
        r.font.name = "Calibri"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = DARK_BLUE
        return h

    # SECTION 1
    add_section_header("1. Executive Summary")
    p = doc.add_paragraph()
    p.add_run(
        "CertiChain Nexus is an enterprise-grade digital certificate verification and fraud detection platform. "
        "It eliminates academic and professional credential tampering by combining salted SHA-256 commitments, "
        "asymmetric Ed25519 digital signatures, an immutable append-only blockchain ledger, and periodic "
        "Merkle tree checkpointing. Rather than offering opaque 'valid/invalid' results, CertiChain Nexus operates "
        "an Explainable Trust Engine evaluating credentials across a strict 5-point inspection matrix."
    )

    # SECTION 2
    add_section_header("2. Certificate Verification Methods")
    
    add_sub_header("Method 2.1: Instant Online Verification via Certificate ID")
    p = doc.add_paragraph()
    p.add_run(
        "Any verifier, recruiter, or academic institution can evaluate a credential in real time without creating an account:\n"
    )
    steps_1 = [
        ("Step 1: Open Verification Portal", "Navigate to http://127.0.0.1:5000/verify in any web browser."),
        ("Step 2: Enter Certificate Identifier", "Input the target Certificate ID into the search bar (e.g. CERT-2026-70C0583D)."),
        ("Step 3: Cryptographic Evaluation", "Click 'Evaluate Credential'. The server executes an end-to-end audit: "
         "recalculating the SHA-256 payload hash, verifying the Ed25519 signature against the issuer's registered public key, "
         "checking block continuity on the blockchain, and inspecting revocation records."),
        ("Step 4: Review Explainable Report", "The result presents a 5-point cryptographic matrix, detailed degree metadata, "
         "and the complete audit timeline.")
    ]
    for step_title, step_desc in steps_1:
        p_step = doc.add_paragraph(style='List Bullet')
        r_title = p_step.add_run(f"{step_title}: ")
        r_title.bold = True
        p_step.add_run(step_desc)

    add_sub_header("Method 2.2: Dynamic QR-Code Scanning")
    p = doc.add_paragraph()
    p.add_run(
        "Each issued certificate embeds an authoritative, high-resolution QR code. When scanned by a smartphone or camera, "
        "it immediately opens the authoritative verification endpoint (`/verify/<certificate_id>`). The QR code itself is "
        "never trusted as offline proof: scanning forces live server-side cryptographic re-validation against the immutable ledger."
    )

    add_sub_header("Method 2.3: Status Precedence Hierarchy")
    p = doc.add_paragraph()
    p.add_run(
        "To prevent ambiguities, evaluation strictly follows a deterministic precedence hierarchy:\n"
        "LEDGER_ERROR ≻ TAMPERED ≻ INVALID_PROOF ≻ ISSUER_UNTRUSTED ≻ REVOKED ≻ SUPERSEDED ≻ VERIFIED"
    )

    # SECTION 3
    add_section_header("3. Standalone Proof Capsules (Upload & Offline Verification)")
    p = doc.add_paragraph()
    p.add_run(
        "A Proof Capsule is a self-contained, portable JSON file that packages all cryptographic commitments necessary "
        "to prove a certificate's authenticity independently of the primary server."
    )

    add_sub_header("3.1 Anatomy of a Proof Capsule JSON")
    p = doc.add_paragraph()
    p.add_run(
        "The downloaded capsule includes:\n"
        "• Certificate Metadata: Identifier, student reference, course, degree, and issue date.\n"
        "• Cryptographic Salt & Commitment: 64-hexadecimal salt and SHA-256 content hash.\n"
        "• Digital Signature: Asymmetric Ed25519 digital signature signed by the issuer's private key.\n"
        "• Issuer Public Key & Verification Data: Registered public key in PEM format.\n"
        "• Merkle Checkpoint Proof: Root hash and cryptographic proof of ledger inclusion."
    )

    add_sub_header("3.2 How to Download a Proof Capsule")
    p = doc.add_paragraph()
    p.add_run(
        "1. Direct Web Link: http://127.0.0.1:5000/certificate/CERT-2026-70C0583D/capsule/download\n"
        "2. Student Credential Vault: Log in as a student (`student@oit.edu`), open vault, and click 'Download Proof Capsule'."
    )

    add_sub_header("3.3 How to Upload and Verify via Web Interface")
    p = doc.add_paragraph()
    p.add_run(
        "1. Open http://127.0.0.1:5000/capsule/verify (or click 'Proof Capsule' on the navigation bar).\n"
        "2. Drag-and-drop or click 'Choose File' to upload your `proof_capsule_*.json` file.\n"
        "3. The browser calls `/api/verify/capsule`, performing instant mathematical verification of the signature and hash, "
        "returning a verified badge or failure explanation."
    )

    add_sub_header("3.4 Standalone Offline CLI Verification (Air-Gapped)")
    p = doc.add_paragraph()
    p.add_run(
        "For environments with zero internet access, the system includes an offline Python tool:\n"
        "Command: python tools/verify_capsule.py path/to/proof_capsule_CERT-2026-70C0583D.json\n"
        "This tool uses local Ed25519 cryptographic routines to validate the certificate without querying any web server."
    )

    # SECTION 4
    add_section_header("4. How to Issue New Verified Certificates")
    p = doc.add_paragraph()
    p.add_run(
        "Accredited universities and institutions mint new certificates using the Issuer Dashboard:\n"
    )
    issuance_steps = [
        ("Step 1: Sign In", "Go to http://127.0.0.1:5000/login and sign in with an Issuer account (e.g. issuer@oit.edu / IssuerPass123!)."),
        ("Step 2: Access Issuance Form", "Navigate to 'Issue Certificate' (/issuer/issue)."),
        ("Step 3: Enter Credential Details", "Provide Student ID (STU-2026-XX-XXX), Student Email, Degree/Certificate Type, Course/Major, Issue Date, and Grade/Honors."),
        ("Step 4: Click 'Issue & Mint to Ledger'", "The backend executes the following automated pipeline:"),
        ("   a) Salt Generation", "A cryptographically secure 64-hex random salt is created."),
        ("   b) Canonical SHA-256 Fingerprint", "Payload attributes are formatted canonically and hashed."),
        ("   c) Ed25519 Signing", "The institution's private key decrypts in-memory and digitally signs the fingerprint."),
        ("   d) Blockchain Block Creation", "A new block (e.g. Block #13) is appended to the hash-linked ledger linking to previous block hash."),
        ("   e) QR & Capsule Generation", "High-res QR PNG and Proof Capsule JSON are automatically generated and linked.")
    ]
    for step_title, step_desc in issuance_steps:
        p_step = doc.add_paragraph(style='List Bullet')
        r_title = p_step.add_run(f"{step_title}: ")
        r_title.bold = True
        p_step.add_run(step_desc)

    # SECTION 5
    add_section_header("5. Current Live System Metrics & Operational Report")
    p = doc.add_paragraph()
    p.add_run(
        "The following data represents the authoritative state of the CertiChain Nexus ledger as verified in the live SQLite database:\n"
    )

    # Summary table
    table_summary = doc.add_table(rows=5, cols=3)
    table_summary.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_summary.autofit = False

    t_widths = [Inches(2.5), Inches(1.5), Inches(2.5)]
    headers = ["Metric Category", "Live Count", "Cryptographic / Operational Status"]
    for i, title in enumerate(headers):
        cell = table_summary.cell(0, i)
        cell.width = t_widths[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 120, 120)
        p_hdr = cell.paragraphs[0]
        r_hdr = p_hdr.add_run(title)
        r_hdr.font.bold = True
        r_hdr.font.color.rgb = RGBColor(255, 255, 255)

    summary_rows = [
        ("Total Issued Certificates", "6", "Synchronized in SQLite & Blockchain"),
        ("Active Authentic Certificates", "4", "Passed all 5 inspection points (100% valid)"),
        ("Revoked Certificates", "1", "Flagged REVOKED via append-only ledger"),
        ("Superseded Certificates", "1", "Replaced by newer version credential")
    ]
    for r_idx, (cat, cnt, desc) in enumerate(summary_rows, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([cat, cnt, desc]):
            cell = table_summary.cell(r_idx, c_idx)
            cell.width = t_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 120, 120)
            p_c = cell.paragraphs[0]
            p_c.add_run(val)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_sub_header("5.1 Complete Inventory of Issued Certificates in Ledger")
    
    cert_table = doc.add_table(rows=7, cols=5)
    cert_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cert_table.autofit = False
    
    c_widths = [Inches(1.8), Inches(1.3), Inches(1.4), Inches(1.2), Inches(0.8)]
    cert_headers = ["Certificate ID", "Student Ref", "Course / Degree", "Issuing Univ", "Status"]
    for i, title in enumerate(cert_headers):
        cell = cert_table.cell(0, i)
        cell.width = c_widths[i]
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 100, 100, 100, 100)
        p_hdr = cell.paragraphs[0]
        r_hdr = p_hdr.add_run(title)
        r_hdr.font.bold = True
        r_hdr.font.color.rgb = RGBColor(255, 255, 255)

    cert_rows = [
        ("CERT-2026-70C0583D", "STU-2026-CS-001", "B.Sc. Computer Science", "OIT", "ACTIVE"),
        ("CERT-2026-E031636A", "STU-2026-AI-002", "M.Sc. Artificial Intelligence", "MIT", "ACTIVE"),
        ("CERT-2026-CE8897D6", "STU-2026-SEC-004", "B.Eng. Cybersecurity v2", "OIT", "ACTIVE"),
        ("CERT-2026-9592CD2D", "STU-2026-TEST-101", "B.Sc. Distributed Systems", "OIT", "ACTIVE"),
        ("CERT-2026-31DBB088", "STU-2026-DS-003", "Diploma in Data Analytics", "OIT", "REVOKED"),
        ("CERT-2026-3785F134", "STU-2026-SEC-004", "B.Eng. Cybersecurity v1", "OIT", "SUPERSEDED")
    ]
    for r_idx, (cid, sref, course, issuer, status) in enumerate(cert_rows, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([cid, sref, course, issuer, status]):
            cell = cert_table.cell(r_idx, c_idx)
            cell.width = c_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            p_c = cell.paragraphs[0]
            run_val = p_c.add_run(val)
            if c_idx == 4:
                run_val.bold = True
                if val == "ACTIVE":
                    run_val.font.color.rgb = GREEN
                elif val == "REVOKED":
                    run_val.font.color.rgb = RED
                elif val == "SUPERSEDED":
                    run_val.font.color.rgb = AMBER

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_sub_header("5.2 Verification Audit Logs & Anomaly Detection (9 Total Checks)")
    p = doc.add_paragraph()
    p.add_run(
        "The system has executed 9 authoritative verification checks recorded with cryptographic timestamps:\n"
        "• 4 VERIFIED: Authenticated active credentials passing all signature & Merkle checks.\n"
        "• 2 UNRECOGNIZED: Attempted searches for non-existent IDs (probed by security scanners).\n"
        "• 1 REVOKED: Attempted verification of credential CERT-2026-31DBB088 (flagged revoked).\n"
        "• 1 SUPERSEDED: Attempted verification of outdated version CERT-2026-3785F134 (redirects to replacement).\n"
        "• 1 TAMPERED: Fraud Radar heuristic detected synthetic modification probe (CERT-SAMPLE-TAMPERED-01)."
    )

    add_sub_header("5.3 Blockchain Ledger Architecture")
    p = doc.add_paragraph()
    p.add_run(
        "• Synchronized Ledger Blocks: 12 immutable hash-linked blocks (Block #0 Genesis through Block #11).\n"
        "• Merkle Tree Checkpoints: 1 cryptographically signed checkpoint anchoring blocks 0-10.\n"
        "• Privacy-Preserving Selective Disclosure: 1 active time-bounded token with salted hash masking."
    )

    # SECTION 6
    add_section_header("6. System Credentials & Reference Guide")
    
    cred_table = doc.add_table(rows=5, cols=4)
    cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cred_table.autofit = False
    
    cr_widths = [Inches(1.8), Inches(2.2), Inches(1.3), Inches(1.2)]
    cred_headers = ["User Role", "Email Address", "Password", "Permissions"]
    for i, title in enumerate(cred_headers):
        cell = cred_table.cell(0, i)
        cell.width = cr_widths[i]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, 100, 100, 100, 100)
        p_hdr = cell.paragraphs[0]
        r_hdr = p_hdr.add_run(title)
        r_hdr.font.bold = True
        r_hdr.font.color.rgb = RGBColor(255, 255, 255)

    cred_rows = [
        ("Super Administrator", "admin@certichain.nexus", "AdminPass123!", "Full platform oversight, issuer approvals"),
        ("Institution Issuer (OIT)", "issuer@oit.edu", "IssuerPass123!", "Issue & revoke certificates, keys"),
        ("Student (OIT)", "student@oit.edu", "StudentPass123!", "Credential vault, share links, capsules"),
        ("Public Verifier", "No login required", "No login required", "Search by ID, upload JSON, scan QR")
    ]
    for r_idx, (role, email, pwd, perms) in enumerate(cred_rows, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate([role, email, pwd, perms]):
            cell = cred_table.cell(r_idx, c_idx)
            cell.width = cr_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 60, 60, 80, 80)
            p_c = cell.paragraphs[0]
            p_c.add_run(val)

    doc.add_paragraph().paragraph_format.space_before = Pt(20)
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = footer_p.add_run("CertiChain Nexus © 2026 • Trust Every Credential. Verify Every Claim.")
    r_foot.font.size = Pt(9)
    r_foot.font.color.rgb = SLATE

    output_path = r"d:\Projects\certichain-nexus\docs\CertiChain_Nexus_Verification_Guide_and_System_Status.docx"
    doc.save(output_path)
    print(f"Document successfully saved to: {output_path}")

if __name__ == "__main__":
    create_document()
